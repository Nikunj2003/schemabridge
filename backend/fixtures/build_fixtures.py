"""Generate the synthetic sample files.

Every row here is invented. The data is shaped to exercise each decision the
engine has to make — safe automatic fixes, genuine ambiguity, identity
conflicts, and delivery failures — so the demo is reproducible rather than
anecdotal.
"""

from __future__ import annotations

import csv
import io
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook

OUT = Path(__file__).resolve().parent / "samples"


def to_csv(rows: list[list[str]]) -> str:
    buffer = io.StringIO(newline="")
    writer = csv.writer(buffer, lineterminator="\n")
    writer.writerows(rows)
    return buffer.getvalue()


# ---------------------------------------------------------------------------
# File 1: legacy export. Terse headers, messy values.
# ---------------------------------------------------------------------------
LEGACY = to_csv(
    [
        ["emp_id", "emp_nm", "email", "doj", "dept", "emp_type"],
        # Whitespace and mixed casing, and a month-name date: all safely fixable.
        [
            "E-1001",
            "  Priya Sharma ",
            "Priya.Sharma@EXAMPLE.COM",
            "14 March 2026",
            "Engineering",
            "Full Time",
        ],
        ["E-1002", "Rahul Verma", "rahul.verma@example.com", "2026-01-06", "Finance", "full-time"],
        ["E-1003", "Asha Rao", "asha.rao@example.com", "2 April 2026", "Operations", "Contract"],
        # Ambiguous date: 03/04 could be 3 April or 4 March. Must be escalated.
        [
            "E-1004",
            "Vikram Nair",
            "vikram.nair@example.com",
            "03/04/2026",
            "Engineering",
            "Full Time",
        ],
        # Unrepairable email: still invalid after safe cleanup, so it fails twice.
        ["E-1005", "Meera Joshi", "meera.joshi@@example..com", "2026-02-17", "Design", "Part Time"],
        # Leading zeros must survive as an opaque identifier.
        ["000123", "Sanjay Gupta", "sanjay.gupta@example.com", "2026-05-11", "Finance", "Intern"],
        # Delivery: the stub fails this one's first attempt, then succeeds on retry.
        ["E-1007", "Deepa Iyer", "deepa.iyer@example.com", "2026-03-02", "Operations", "Full Time"],
        # Delivery: the stub rejects this one permanently.
        [
            "E-1008",
            "Arjun Mehta",
            "arjun.mehta@example.com",
            "2026-04-20",
            "Engineering",
            "Contract",
        ],
    ]
)

# ---------------------------------------------------------------------------
# File 2: newer export. Different header spellings, a merge, a conflict, and one
# genuinely ambiguous column.
# ---------------------------------------------------------------------------
HR_EXPORT = to_csv(
    [
        [
            "Employee Number",
            "Full Name",
            "Work Email",
            # Ambiguous: could be the start date or the end date.
            "Date",
            "Joining Date",
            "Division",
            "Engagement Type",
        ],
        # Same person as E-1002. Values agree once spellings are canonicalised
        # ("Permanent" means the same as "full-time"), so this merges silently.
        [
            "E-1002",
            "Rahul Verma",
            "rahul.verma@example.com",
            "",
            "2026-01-06",
            "Finance",
            "Permanent",
        ],
        # Same person as E-1003, but this file asserts a different joining date.
        # Both values are valid; only a human can say which is right.
        ["E-1003", "Asha Rao", "asha.rao@example.com", "", "2024-11-30", "Operations", "Contract"],
        [
            "E-2001",
            "Nisha Kapoor",
            "nisha.kapoor@example.com",
            "",
            "2026-06-01",
            "Marketing",
            "Full Time",
        ],
        [
            "E-2002",
            "Karan Singh",
            "karan.singh@example.com",
            "2026-08-31",
            "2026-06-15",
            "Marketing",
            "Part Time",
        ],
    ]
)

# ---------------------------------------------------------------------------
# File 3: clean file. Should complete with no human input at all.
# ---------------------------------------------------------------------------
CLEAN = to_csv(
    [
        ["employeeId", "fullName", "workEmail", "startDate", "department", "employmentType"],
        ["C-9001", "Leena Das", "leena.das@example.com", "2026-02-02", "Support", "full_time"],
        ["C-9002", "Omar Faruq", "omar.faruq@example.com", "2026-02-09", "Support", "part_time"],
        ["C-9003", "Tara Bose", "tara.bose@example.com", "2026-03-16", "Support", "contract"],
    ]
)


def build_workbook() -> bytes:
    """An Excel export with real typed dates and one header no alias table knows."""
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Staff"

    sheet.append(
        [
            "Staff ID",
            "Name",
            "Company Email",
            "Joining Date",
            "Team",
            # No alias table knows this. It is what model assistance is for.
            "Cost Centre Ref",
        ]
    )

    rows: list[list[object]] = [
        [
            "E-3001",
            "Ananya Bhat",
            "ananya.bhat@example.com",
            datetime(2026, 1, 20),
            "Platform",
            "CC-4410",
        ],
        [
            "E-3002",
            "Rohan Pillai",
            "rohan.pillai@example.com",
            datetime(2026, 4, 6),
            "Platform",
            "CC-4410",
        ],
        # Same person as E-1001 in the legacy file; the values agree, so it merges.
        [
            "E-1001",
            "Priya Sharma",
            "priya.sharma@example.com",
            datetime(2026, 3, 14),
            "Engineering",
            "CC-2201",
        ],
    ]
    for row in rows:
        added = sheet.append(row)  # type: ignore[func-returns-value]
        del added
        sheet.cell(row=sheet.max_row, column=4).number_format = "yyyy-mm-dd"

    for column in "ABCDEF":
        sheet.column_dimensions[column].width = 22

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


FILES: dict[str, str] = {
    "employees-legacy.csv": LEGACY,
    "employees-hr-export.csv": HR_EXPORT,
    "employees-clean.csv": CLEAN,
}

DESCRIPTIONS = {
    "employees-legacy.csv": "messy values, ambiguous date, unrepairable email",
    "employees-hr-export.csv": "different headers, a merge, a conflict, an ambiguous column",
    "employees-clean.csv": "requires no human input",
    "employees-directory.xlsx": "typed Excel dates, one unknown header",
}


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    for name, content in FILES.items():
        (OUT / name).write_text(content, encoding="utf-8")
    (OUT / "employees-directory.xlsx").write_bytes(build_workbook())

    print(f"Wrote fixtures to {OUT}")
    for name, description in DESCRIPTIONS.items():
        print(f"  {name:<28} {description}")


if __name__ == "__main__":
    main()
