"""Uploaded files are untrusted. Every limit is enforced server-side."""

from __future__ import annotations

import io

from openpyxl import Workbook

from schemabridge.domain.target import ValueKind
from schemabridge.ingest.csv_source import parse_csv
from schemabridge.ingest.limits import INGEST_LIMITS
from schemabridge.ingest.profile import profile_columns
from schemabridge.ingest.xlsx_source import parse_xlsx


def ok(result: object) -> object:
    assert getattr(result, "ok", False), getattr(result, "error", result)
    return result


class TestParseCsv:
    def test_reads_headers_and_rows_with_stable_ids(self) -> None:
        result = parse_csv("f1", "a.csv", "emp_id,name\nE-1,Asha\nE-2,Rahul\n")
        assert result.ok
        assert [c.header for c in result.file.columns] == ["emp_id", "name"]
        assert len(result.rows) == 2
        first = result.rows[0]
        assert first.values[result.file.columns[0].id] == "E-1"
        assert first.row == 2  # header is row 1

    def test_preserves_leading_zeros(self) -> None:
        result = parse_csv("f1", "a.csv", "emp_id\n000123\n")
        assert result.ok
        assert result.rows[0].values[result.file.columns[0].id] == "000123"

    def test_handles_quoted_commas_and_newlines(self) -> None:
        result = parse_csv("f1", "a.csv", 'name,note\n"Smith, John","line1\nline2"\n')
        assert result.ok
        cols = result.file.columns
        assert result.rows[0].values[cols[0].id] == "Smith, John"
        assert "line1" in result.rows[0].values[cols[1].id]

    def test_strips_a_byte_order_mark(self) -> None:
        result = parse_csv("f1", "a.csv", "﻿emp_id,name\nE-1,Asha\n")
        assert result.ok
        assert result.file.columns[0].normalized_header == "empid"

    def test_keeps_duplicate_headers_distinct(self) -> None:
        result = parse_csv("f1", "a.csv", "date,date\n2026-01-01,2026-02-02\n")
        assert result.ok
        assert len(result.file.columns) == 2
        c1, c2 = result.file.columns
        assert c1.id != c2.id
        assert result.rows[0].values[c1.id] == "2026-01-01"
        assert result.rows[0].values[c2.id] == "2026-02-02"

    def test_skips_blank_lines(self) -> None:
        result = parse_csv("f1", "a.csv", "emp_id\nE-1\n\n   \nE-2\n")
        assert result.ok
        assert len(result.rows) == 2

    def test_rejects_a_file_with_no_data_rows(self) -> None:
        result = parse_csv("f1", "a.csv", "emp_id,name\n")
        assert not result.ok
        assert "no data rows" in result.error.lower()

    def test_rejects_an_empty_header_row(self) -> None:
        assert not parse_csv("f1", "a.csv", "\n\n").ok

    def test_rejects_too_many_columns(self) -> None:
        header = ",".join(f"c{i}" for i in range(INGEST_LIMITS.max_columns + 1))
        result = parse_csv("f1", "a.csv", f"{header}\n1\n")
        assert not result.ok
        assert "column" in result.error.lower()

    def test_rejects_too_many_rows(self) -> None:
        rows = "\n".join(f"E-{i}" for i in range(INGEST_LIMITS.max_rows_per_file + 1))
        result = parse_csv("f1", "a.csv", f"emp_id\n{rows}\n")
        assert not result.ok
        assert "row" in result.error.lower()

    def test_rejects_an_oversized_cell(self) -> None:
        big = "x" * (INGEST_LIMITS.max_cell_length + 1)
        result = parse_csv("f1", "a.csv", f"note\n{big}\n")
        assert not result.ok
        assert "too long" in result.error.lower()


class TestProfileColumns:
    def test_summarises_and_caps_samples(self) -> None:
        rows = "\n".join(f"E-{i},u{i}@x.com" for i in range(1, 7))
        parsed = parse_csv("f1", "a.csv", f"emp_id,email\n{rows}\n")
        assert parsed.ok
        profiles = profile_columns(parsed.file, parsed.rows)
        assert len(profiles) == 2
        email = profiles[1]
        assert email.non_empty_count == 6
        assert email.unique_ratio == 1
        assert ValueKind.EMAIL in email.detected_kinds
        assert len(email.samples) <= INGEST_LIMITS.max_samples_per_column

    def test_a_wholly_blank_line_is_not_a_row(self) -> None:
        """An empty line in a CSV is formatting, not a record with no values."""
        parsed = parse_csv("f1", "a.csv", "dept\nFinance\n\nOps\n")
        assert parsed.ok
        dept = profile_columns(parsed.file, parsed.rows)[0]
        assert (dept.total_count, dept.non_empty_count) == (2, 2)

    def test_a_barely_filled_column_reports_as_such(self) -> None:
        """`total_count` spans every row; `non_empty_count` only the filled ones.

        The gap is the signal. Reporting both as the filled count would tell the
        model a column with one value in ten rows was complete — which is exactly
        the evidence that should make it doubt a confident-looking header.
        """
        rows = "\n".join([f"E-{index},," for index in range(1, 10)] + ["E-10,,x"])
        parsed = parse_csv("f1", "a.csv", f"id,note,flag\n{rows}\n")
        assert parsed.ok
        note, flag = profile_columns(parsed.file, parsed.rows)[1:]
        assert (note.total_count, note.non_empty_count) == (10, 0)
        assert (flag.total_count, flag.non_empty_count) == (10, 1)


def workbook_bytes(build: object) -> bytes:
    wb = Workbook()
    sheet = wb.active
    assert sheet is not None
    assert callable(build)
    build(sheet, wb)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


class TestParseXlsx:
    def test_reads_a_single_populated_sheet(self) -> None:
        def build(sheet: object, _wb: object) -> None:
            sheet.append(["employee_id", "full_name"])  # type: ignore[attr-defined]
            sheet.append(["E-1", "Asha Rao"])  # type: ignore[attr-defined]
            sheet.append(["E-2", "Rahul Verma"])  # type: ignore[attr-defined]

        result = parse_xlsx("f1", "hr.xlsx", workbook_bytes(build))
        assert result.ok, result.error
        assert [c.header for c in result.file.columns] == ["employee_id", "full_name"]
        assert len(result.rows) == 2

    def test_converts_a_typed_date_to_a_calendar_date(self) -> None:
        from datetime import datetime

        def build(sheet: object, _wb: object) -> None:
            sheet.append(["start_date"])  # type: ignore[attr-defined]
            sheet.append([datetime(2026, 3, 14)])  # type: ignore[attr-defined]

        result = parse_xlsx("f1", "hr.xlsx", workbook_bytes(build))
        assert result.ok, result.error
        assert result.rows[0].values[result.file.columns[0].id] == "2026-03-14"

    def test_rejects_a_formula_cell(self) -> None:
        def build(sheet: object, _wb: object) -> None:
            sheet.append(["employee_id", "derived"])  # type: ignore[attr-defined]
            sheet.append(["E-1", '=A2&"-x"'])  # type: ignore[attr-defined]

        result = parse_xlsx("f1", "hr.xlsx", workbook_bytes(build))
        assert not result.ok
        assert "formula" in result.error.lower()

    def test_refuses_two_populated_sheets(self) -> None:
        wb = Workbook()
        first = wb.active
        assert first is not None
        first.title = "Employees"
        first.append(["employee_id"])
        first.append(["E-1"])
        second = wb.create_sheet("Contractors")
        second.append(["employee_id"])
        second.append(["C-1"])
        buffer = io.BytesIO()
        wb.save(buffer)

        result = parse_xlsx("f1", "hr.xlsx", buffer.getvalue())
        assert not result.ok
        assert "worksheet" in result.error.lower()

    def test_skips_blank_spacer_rows(self) -> None:
        def build(sheet: object, _wb: object) -> None:
            sheet.append(["employee_id"])  # type: ignore[attr-defined]
            sheet.append(["E-1"])  # type: ignore[attr-defined]
            sheet.append([None])  # type: ignore[attr-defined]
            sheet.append(["E-2"])  # type: ignore[attr-defined]

        result = parse_xlsx("f1", "hr.xlsx", workbook_bytes(build))
        assert result.ok, result.error
        assert len(result.rows) == 2

    def test_reports_headers_without_data(self) -> None:
        def build(sheet: object, _wb: object) -> None:
            sheet.append(["employee_id", "full_name"])  # type: ignore[attr-defined]

        assert not parse_xlsx("f1", "hr.xlsx", workbook_bytes(build)).ok

    def test_rejects_a_file_that_is_not_a_workbook(self) -> None:
        result = parse_xlsx("f1", "hr.xlsx", b"not,a,workbook\n1,2,3")
        assert not result.ok
        assert "could not be read" in result.error.lower()

    def test_a_formula_without_a_cached_value_is_refused_not_blanked(self) -> None:
        """The subtle case: a programmatically written workbook has no cache.

        Reading only cached values turns such a formula into an empty string,
        so the column looks present but is quietly empty. Refusing the file is
        the honest outcome.
        """

        def build(sheet: object, _wb: object) -> None:
            sheet.append(["employee_id", "start_date"])  # type: ignore[attr-defined]
            sheet.append(["E-1", "=TODAY()"])  # type: ignore[attr-defined]

        result = parse_xlsx("f1", "hr.xlsx", workbook_bytes(build))
        assert not result.ok
        assert "formula" in result.error.lower()
        # The reviewer is told exactly which cell, not just that something failed.
        assert "B2" in result.error

    def test_accepts_a_workbook_with_no_formulas(self) -> None:
        def build(sheet: object, _wb: object) -> None:
            sheet.append(["employee_id", "note"])  # type: ignore[attr-defined]
            sheet.append(["E-1", "plain text, no formula"])  # type: ignore[attr-defined]

        result = parse_xlsx("f1", "hr.xlsx", workbook_bytes(build))
        assert result.ok, result.error
        assert result.rows[0].values[result.file.columns[1].id] == "plain text, no formula"
